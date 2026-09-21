#!/usr/bin/env python3
"""
Mentee Mapping SOP — poster system
Estimate → Identity. Unique composition per beat. Motion with meaning.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Cool stone + forest + signal orange (no purple template, no cream)
INK = RGBColor(0x0C, 0x0E, 0x0D)
STONE = RGBColor(0xE8, 0xEC, 0xE9)
MUTE = RGBColor(0x6E, 0x78, 0x72)
TEAL = RGBColor(0x24, 0x5C, 0x4E)
ORANGE = RGBColor(0xFD, 0x4F, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC8, 0xD0, 0xCB)
CHAR = RGBColor(0x12, 0x16, 0x14)

FONT_D = "Source Serif 4"
FONT_B = "DM Sans"
FONT_M = "JetBrains Mono"

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 14
ASSETS = Path("/workspace/docs/assets")

_anim_id = {"n": 10}


def _fid():
    _anim_id["n"] += 1
    return _anim_id["n"]


def _set_font(run, size, bold, color, name):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def tb(slide, left, top, width, height, text, *, size=16, bold=False,
       color=INK, align=PP_ALIGN.LEFT, font=FONT_B):
    sh = slide.shapes.add_textbox(left, top, width, height)
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_font(r, size, bold, color, font)
    return sh


def ml(slide, left, top, width, height, lines, *, size=15, color=INK,
       font=FONT_B, before=12, bullet=False):
    sh = slide.shapes.add_textbox(left, top, width, height)
    tf = sh.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(opts.get("before", before if i else 0))
        r = p.add_run()
        r.text = (("·  " if bullet or opts.get("bullet") else "") + text)
        _set_font(
            r,
            opts.get("size", size),
            opts.get("bold", False),
            opts.get("color", color),
            opts.get("font", font),
        )
    return sh


def pic(slide, name, left=0, top=0, width=None, height=None):
    return slide.shapes.add_picture(
        str(ASSETS / name), left, top, width=width or W, height=height or H
    )


def rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def oval(slide, left, top, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def coord(slide, n, color=MUTE):
    return tb(
        slide, Inches(11.7), Inches(7.1), Inches(1.3), Inches(0.28),
        f"{n:02d}/{TOTAL:02d}", size=10, color=color, align=PP_ALIGN.RIGHT, font=FONT_M,
    )


# --- animations ---

def _entr(spid, dur, delay, click, filter_="fade", preset=10):
    node = "clickEffect" if click else "withEffect"
    d = "indefinite" if click else str(delay)
    return f"""
    <p:par>
      <p:cTn id="{_fid()}" presetID="{preset}" presetClass="entr" presetSubtype="0"
             fill="hold" grpId="0" nodeType="{node}">
        <p:stCondLst><p:cond delay="{d}"/></p:stCondLst>
        <p:childTnLst>
          <p:set>
            <p:cBhvr>
              <p:cTn id="{_fid()}" dur="1" fill="hold">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
              </p:cTn>
              <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
              <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
            </p:cBhvr>
            <p:to><p:strVal val="visible"/></p:to>
          </p:set>
          <p:animEffect transition="in" filter="{filter_}">
            <p:cBhvr>
              <p:cTn id="{_fid()}" dur="{dur}"/>
              <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
            </p:cBhvr>
          </p:animEffect>
        </p:childTnLst>
      </p:cTn>
    </p:par>"""


def animate(slide, sequence, *, stagger=300, dur=650, on_click=False, wipe=False):
    if not sequence:
        return
    _anim_id["n"] = 10
    effects, ids = [], []
    delay = 0
    filt = "wipe(down)" if wipe else "fade"
    preset = 22 if wipe else 10
    for i, item in enumerate(sequence):
        shapes = item if isinstance(item, (list, tuple)) else [item]
        click = bool(on_click and i > 0)
        for sh in shapes:
            if sh is None:
                continue
            ids.append(sh.shape_id)
            effects.append(_entr(sh.shape_id, dur, 0 if click else delay, click, filt, preset))
        if not on_click:
            delay += stagger

    # bldLst
    cSld = slide._element.find(qn("p:cSld"))
    if cSld is not None:
        bld = cSld.find(qn("p:bldLst"))
        if bld is None:
            bld = etree.SubElement(cSld, qn("p:bldLst"))
        have = {el.get("spid") for el in bld.findall(qn("p:bldP"))}
        for sid in ids:
            if str(sid) not in have:
                el = etree.SubElement(bld, qn("p:bldP"))
                el.set("spid", str(sid))
                el.set("grpId", "0")
                el.set("animBg", "1")

    timing = etree.fromstring(f"""
    <p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:tnLst>
        <p:par>
          <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
            <p:childTnLst>
              <p:seq concurrent="1" nextAc="seek">
                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="3" fill="hold">
                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                        <p:childTnLst>
                          <p:par>
                            <p:cTn id="4" fill="hold">
                              <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                              <p:childTnLst>{''.join(effects)}</p:childTnLst>
                            </p:cTn>
                          </p:par>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
                <p:prevCondLst>
                  <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
                </p:prevCondLst>
                <p:nextCondLst>
                  <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
                </p:nextCondLst>
              </p:seq>
            </p:childTnLst>
          </p:cTn>
        </p:par>
      </p:tnLst>
    </p:timing>""")
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    ext = sld.find(qn("p:extLst"))
    (ext.addprevious if ext is not None else sld.append)(timing)


# --- slides ---

def s01(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_title.png")
    a = tb(s, Inches(0.85), Inches(1.85), Inches(7), Inches(0.35),
           "MENTORS  ·  SOP", size=13, bold=True, color=ORANGE, font=FONT_M)
    b = tb(s, Inches(0.85), Inches(2.4), Inches(7.5), Inches(1.5),
           "Estimate\nto identity.", size=52, bold=True, color=WHITE, font=FONT_D)
    c = tb(s, Inches(0.85), Inches(4.3), Inches(7), Inches(0.7),
           "How mentee mapping moved from facility guesses\nto a living Kobo pin for every HCW.",
           size=16, color=SOFT, font=FONT_B)
    d = tb(s, Inches(0.85), Inches(6.45), Inches(6), Inches(0.3),
           "Jacaranda Health", size=13, color=MUTE, font=FONT_B)
    animate(s, [a, b, c, d], stagger=380, dur=750)


def s02(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_journey.png")
    # Journey art carries the story — only a whisper of title
    t = tb(s, Inches(0.55), Inches(0.35), Inches(8), Inches(0.45),
           "Where we came from", size=22, bold=True, color=INK, font=FONT_D)
    animate(s, [t], dur=700)
    coord(s, 2)


def s03(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_facility.png")
    k = tb(s, Inches(0.7), Inches(1.3), Inches(6), Inches(0.3),
           "2018 — 2023", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.75), Inches(6.2), Inches(1.1),
           "We counted\nthe building.", size=36, bold=True, color=INK, font=FONT_D)
    body = ml(s, Inches(0.7), Inches(3.3), Inches(5.8), Inches(2.8), [
        "A CME in a facility meant every HCW there was “trained.”",
        "7,441 HCWs — an estimate from maternity-wing headcount.",
        "Attrition made the number look better than the truth.",
    ], size=15, color=MUTE, bullet=True, before=16)
    animate(s, [k, t, body], stagger=320, dur=600, wipe=True)
    coord(s, 3)


def s04(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_random.png")
    k = tb(s, Inches(0.7), Inches(1.3), Inches(5), Inches(0.3),
           "THE PIVOT", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.75), Inches(5.5), Inches(1.1),
           "We typed\nthe person.", size=36, bold=True, color=INK, font=FONT_D)
    body = ml(s, Inches(0.7), Inches(3.3), Inches(5.3), Inches(2.8), [
        "Facility → name → phone (mentee_id) → activities.",
        "Better than assumptions. Still memory-bound.",
        "Mistyped phones broke the unique key.",
    ], size=15, color=MUTE, bullet=True, before=16)
    animate(s, [k, t, body], stagger=320, dur=600)
    coord(s, 4)


def s05(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_poled.png")
    k = tb(s, Inches(0.7), Inches(1.15), Inches(5), Inches(0.3),
           "CURRENT", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.6), Inches(5.3), Inches(1.0),
           "We preload\nthe pin.", size=36, bold=True, color=INK, font=FONT_D)
    cols = [
        ("Map", "POs register eligible HCWs on site."),
        ("Select", "IFMs pick names in Kobo — no inventing."),
        ("Link", "WhatsApp joins in-person to DELTA."),
    ]
    shapes = [k, t]
    for i, (h, b) in enumerate(cols):
        y = Inches(3.0) + i * Inches(0.95)
        hh = tb(s, Inches(0.7), y, Inches(5), Inches(0.35),
                h, size=16, bold=True, color=TEAL, font=FONT_B)
        bb = tb(s, Inches(0.7), y + Inches(0.35), Inches(5), Inches(0.45),
                b, size=14, color=MUTE, font=FONT_B)
        shapes.append([hh, bb])
    v = tb(s, Inches(0.7), Inches(6.1), Inches(5.2), Inches(0.6),
           "Same mentee. In person, hybrid, or DELTA.",
           size=14, bold=True, color=INK, font=FONT_B)
    shapes.append(v)
    animate(s, shapes, stagger=260, dur=550)
    coord(s, 5)


def s06(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_problem.png")
    t = tb(s, Inches(0.7), Inches(0.55), Inches(12), Inches(0.7),
           "What broke the old way", size=34, bold=True, color=INK, font=FONT_D)
    items = [
        ("01", "No mentee-level truth", "Facilities can’t tell you who was mentored."),
        ("02", "Field noise", "Typed phones and names drifted under pressure."),
        ("03", "Manual preload pain", "Quality rose. Keeping tools current did not."),
        ("04", "Attrition lag", "Inactive people haunted the forms."),
    ]
    shapes = [t]
    for i, (n, h, b) in enumerate(items):
        y = Inches(1.6) + i * Inches(1.2)
        nn = tb(s, Inches(0.7), y, Inches(1.0), Inches(0.5),
                n, size=24, bold=True, color=ORANGE, font=FONT_M)
        hh = tb(s, Inches(1.9), y, Inches(10), Inches(0.4),
                h, size=20, bold=True, color=INK, font=FONT_D)
        bb = tb(s, Inches(1.9), y + Inches(0.45), Inches(10), Inches(0.4),
                b, size=15, color=MUTE, font=FONT_B)
        shapes.append([nn, hh, bb])
    animate(s, shapes, stagger=280, dur=500, on_click=True)
    coord(s, 6)


def s07(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_pipeline.png")
    title = tb(s, Inches(0.7), Inches(0.45), Inches(12), Inches(0.55),
               "How the tools stay true", size=28, bold=True, color=INK, font=FONT_D)
    subtitle = tb(s, Inches(0.7), Inches(1.05), Inches(12), Inches(0.35),
                  "Database → fragments → builders → proof → field.",
                  size=15, color=MUTE, font=FONT_B)
    legend = ml(
        s, Inches(0.7), Inches(5.55), Inches(12), Inches(1.3),
        [
            ("01 Sync — clean source of truth", {"bold": True, "size": 13, "color": INK}),
            ("02 Blocks · 03 Build · 04 Check · 05 Ship — one grammar across tools",
             {"size": 13, "color": MUTE}),
        ],
        before=8,
    )
    animate(s, [title, subtitle, legend], stagger=350, dur=600)
    coord(s, 7)



def s08(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_mist.png")
    oval(s, Inches(0.7), Inches(0.55), Inches(0.14), Inches(0.14), ORANGE)
    k = tb(s, Inches(1.0), Inches(0.48), Inches(6), Inches(0.3),
           "SYNC", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.7),
           "One sheet. No ghosts.", size=34, bold=True, color=INK, font=FONT_D)
    rows = [
        ("Pull", "Latest mentee + IFM rows into the working sheet."),
        ("Clean", "Names, IDs, counties, facilities — Kobo-safe."),
        ("Normalize", "Active / Inactive · MENTORS / Newborn / Both."),
        ("Drop", "Missing ID, name, county, facility, or code — gone."),
        ("Dedupe", "Facility–program–status before any builder runs."),
    ]
    shapes = [k, t]
    for i, (h, b) in enumerate(rows):
        y = Inches(2.0) + i * Inches(0.85)
        hh = tb(s, Inches(0.7), y, Inches(2.2), Inches(0.4),
                h, size=16, bold=True, color=TEAL, font=FONT_B)
        bb = tb(s, Inches(3.1), y, Inches(9.5), Inches(0.5),
                b, size=15, color=INK, font=FONT_B)
        shapes.append([hh, bb])
    animate(s, shapes, stagger=220, dur=480)
    coord(s, 8)


def s09(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_mist.png")
    oval(s, Inches(0.7), Inches(0.55), Inches(0.14), Inches(0.14), ORANGE)
    k = tb(s, Inches(1.0), Inches(0.48), Inches(6), Inches(0.3),
           "BLOCKS", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.65),
           "Shared grammar.", size=34, bold=True, color=INK, font=FONT_D)
    groups = [
        ("Survey", [
            "type · name · label · hint",
            "required / messages",
            "relevant · constraint",
            "calculation · choice_filter",
            "appearance · notes",
        ]),
        ("Choices", ["list_name", "name", "label", "choice_filter"]),
        ("Settings", ["allow_choice_duplicates"]),
    ]
    shapes = [k, t]
    for i, (name, lines) in enumerate(groups):
        x = Inches(0.7) + i * Inches(4.15)
        rule = rect(s, x, Inches(2.0), Inches(0.7), Inches(0.05), ORANGE if i == 0 else TEAL)
        hh = tb(s, x, Inches(2.25), Inches(3.8), Inches(0.4),
                name, size=18, bold=True, color=INK, font=FONT_D)
        bb = ml(s, x, Inches(2.85), Inches(3.8), Inches(3.5),
                lines, size=14, color=MUTE, before=12)
        shapes.extend([rule, hh, bb])
    animate(s, shapes, stagger=280, dur=550)
    coord(s, 9)


def s10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_mist.png")
    oval(s, Inches(0.7), Inches(0.55), Inches(0.14), Inches(0.14), ORANGE)
    k = tb(s, Inches(1.0), Inches(0.48), Inches(6), Inches(0.3),
           "BUILD", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.65),
           "Five tools. One spine.", size=34, bold=True, color=INK, font=FONT_D)
    forms = [
        "EmONC Curriculum Tracking Form",
        "Newborn Curriculum Tracking Form",
        "MoH Skills Assessment Checklist",
        "EmONC Knowledge Assessment",
        "Newborn Knowledge Assessment",
    ]
    shapes = [k, t]
    for i, name in enumerate(forms):
        y = Inches(2.0) + i * Inches(0.85)
        n = tb(s, Inches(0.7), y, Inches(0.9), Inches(0.45),
               f"{i+1:02d}", size=18, bold=True, color=ORANGE, font=FONT_M)
        nm = tb(s, Inches(1.8), y, Inches(10.5), Inches(0.45),
                name, size=18, bold=True, color=INK, font=FONT_D)
        shapes.append([n, nm])
    animate(s, shapes, stagger=240, dur=480)
    coord(s, 10)


def s11(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_mist.png")
    oval(s, Inches(0.7), Inches(0.55), Inches(0.14), Inches(0.14), ORANGE)
    k = tb(s, Inches(1.0), Inches(0.48), Inches(6), Inches(0.3),
           "CHECK", size=12, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.7), Inches(1.0), Inches(12), Inches(0.65),
           "Prove it. Then ship.", size=34, bold=True, color=INK, font=FONT_D)
    checks = [
        "Active + correctly programmed records only",
        "Relevance matches field names",
        "Choices align with select_one / select_multiple",
        "Required + choice filters behave",
        "Messages and hints read clean",
        "Hide logic waits for answers",
        "No duplicate facilities, mentees, or IFMs",
    ]
    shapes = [k, t]
    for i, c in enumerate(checks):
        y = Inches(1.85) + i * Inches(0.65)
        mark = tb(s, Inches(0.7), y, Inches(0.45), Inches(0.35),
                  "→", size=14, bold=True, color=TEAL, font=FONT_B)
        line = tb(s, Inches(1.25), y, Inches(11.2), Inches(0.4),
                  c, size=15, color=INK, font=FONT_B)
        shapes.append([mark, line])
    animate(s, shapes, stagger=180, dur=420)
    coord(s, 11)


def s12(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_missing.png")
    t = tb(s, Inches(0.55), Inches(0.35), Inches(12), Inches(0.5),
           "When a pin doesn’t appear", size=26, bold=True, color=INK, font=FONT_D)
    groups = [
        ("Identity", [
            "Invalid mentee_id (format, spaces, length, prefix)",
            "Inactive / blank status or date_activated",
            "Missing from Mentee or IFM database",
        ]),
        ("Program", [
            "Wrong Program value (Newborn / EmONC / Both)",
            "EmONC In-person = No",
            "Both Newborn flags blank/No → Error!",
        ]),
        ("Facility", [
            "Name typos break relevance (Centre/Center)",
            "Reused facility codes — one dropped",
            "Code blank or mapped to wrong site",
        ]),
    ]
    shapes = [t]
    for i, (name, lines) in enumerate(groups):
        x = Inches(0.55) + i * Inches(4.2)
        # header sits in the colored bar zone of the asset (~280-360px → ~1.45-1.9")
        hh = tb(s, x + Inches(0.25), Inches(1.55), Inches(3.5), Inches(0.4),
                name, size=16, bold=True, color=WHITE, font=FONT_B)
        bb = ml(s, x + Inches(0.25), Inches(2.3), Inches(3.6), Inches(4.0),
                lines, size=13, color=MUTE, bullet=True, before=14)
        shapes.extend([hh, bb])
    animate(s, shapes, stagger=300, dur=550)
    coord(s, 12)


def s13(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_deploy.png")
    k = tb(s, Inches(0.85), Inches(2.5), Inches(8), Inches(0.35),
           "SHIP", size=14, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.85), Inches(3.05), Inches(8), Inches(1.2),
           "Upload.\nVersion.\nRelease.", size=40, bold=True, color=WHITE, font=FONT_D)
    animate(s, [k, t], stagger=420, dur=800)
    coord(s, 13, color=SOFT)


def s14(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    pic(s, "poster_future.png")
    k = tb(s, Inches(0.85), Inches(2.45), Inches(8), Inches(0.35),
           "NEXT", size=14, bold=True, color=ORANGE, font=FONT_M)
    t = tb(s, Inches(0.85), Inches(3.0), Inches(8), Inches(0.8),
           "IFMs place the pin.", size=36, bold=True, color=WHITE, font=FONT_D)
    b = tb(s, Inches(0.85), Inches(4.0), Inches(7.5), Inches(0.7),
           "A dedicated Kobo tool for facility-led mentee mapping —\nwhere care happens.",
           size=16, color=SOFT, font=FONT_B)
    animate(s, [k, t, b], stagger=380, dur=750)
    coord(s, 14, color=SOFT)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14):
        fn(prs)
    out = Path("/workspace/docs/Mentee_Mapping_SOP.pptx")
    prs.save(out)
    print(f"Saved {out} · {len(prs.slides)} slides")
    return out


if __name__ == "__main__":
    build()
